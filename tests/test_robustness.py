"""Robustness against packages this parser has not seen.

Only one export package was available while building this, but the other
integrations will contain transformation types, task types and structures that
are not in it. The rule is: never crash, never silently drop an object, and
flag anything not fully understood so a reviewer sees it.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from iics_parser.extract.package import Asset
from iics_parser.model.ir import (Connection, Integration, IntegrationMeta, Mapping,
                                  Step, TaskParameter)
from iics_parser.parse.mapping import _resolve_kind, flow_string, parse_mapping
from iics_parser.render.rows import field_level_rows, mapping_detail_rows


def _write_mapping(tmp_path: Path, name: str, content: dict, class_info: dict) -> Asset:
    """Build a .DTEMPLATE asset directory on disk."""
    root = tmp_path / f"{name}.DTEMPLATE"
    (root / "bin").mkdir(parents=True)
    (root / "bin" / "@3.bin").write_text(
        json.dumps({"content": content, "metadata": {"$$classInfo": class_info}})
    )
    (root / "fileRecord.json").write_text(
        json.dumps([{"id": "@3", "type": "IMFOBJECT", "name": name}])
    )
    return Asset(name=name, obj_type="DTEMPLATE", path=root)


# ---------------------------------------------------------- type resolution

@pytest.mark.parametrize("class_name,expected", [
    ("com.informatica.metadata.template.tx.tmplsource.TmplSource", "Source"),
    ("com.informatica.metadata.template.tx.tmpljoiner.TmplJoiner", "Joiner"),
    ("com.informatica.metadata.template.tx.tmplnormalizer.TmplNormalizer", "Normalizer"),
    ("com.informatica.metadata.template.tx.tmplunion.TmplUnion", "Union"),
])
def test_known_types_resolve(class_name, expected):
    kind, unsupported = _resolve_kind(class_name)
    assert (kind, unsupported) == (expected, False)


def test_unknown_type_is_named_not_dropped():
    """An unmapped transformation still gets a readable name and a flag."""
    kind, unsupported = _resolve_kind(
        "com.informatica.metadata.template.tx.tmplquantum.TmplQuantumFlux"
    )
    assert kind == "QuantumFlux"
    assert unsupported is True


def test_missing_class_info_does_not_crash():
    kind, unsupported = _resolve_kind("")
    assert kind == "Unknown" and unsupported is True


def test_unknown_transformation_is_reported_and_kept(tmp_path):
    asset = _write_mapping(
        tmp_path, "m_FUTURE",
        content={
            "transformations": [
                {"$$ID": 1, "$$class": 7, "name": "src_a"},
                {"$$ID": 2, "$$class": 99, "name": "weird_step"},
                {"$$ID": 3, "$$class": 9, "name": "tgt_b"},
            ],
            "links": [
                {"fromTransformation": {"##ID": 1}, "toTransformation": {"##ID": 2}},
                {"fromTransformation": {"##ID": 2}, "toTransformation": {"##ID": 3}},
            ],
        },
        class_info={
            "7": "com.informatica.metadata.template.tx.tmplsource.TmplSource",
            "9": "com.informatica.metadata.template.tx.tmpltarget.TmplTarget",
            "99": "com.informatica.metadata.template.tx.tmplnew.TmplBrandNew",
        },
    )
    warnings = []
    mapping = parse_mapping(asset, connections={}, warnings=warnings)

    assert [t.name for t in mapping.transformations] == ["src_a", "weird_step", "tgt_b"]
    assert any("weird_step" in w and "BrandNew" in w for w in warnings)
    # It still takes its place in the flow rather than disappearing.
    assert flow_string(mapping) == "src_a->weird_step->tgt_b"


# ----------------------------------------------------------- malformed data

def test_mapping_with_no_graph_warns_but_returns(tmp_path):
    root = tmp_path / "m_EMPTY.DTEMPLATE"
    root.mkdir(parents=True)
    asset = Asset(name="m_EMPTY", obj_type="DTEMPLATE", path=root)
    warnings = []
    mapping = parse_mapping(asset, connections={}, warnings=warnings)
    assert mapping.name == "m_EMPTY"
    assert mapping.transformations == []
    assert any("no mapping graph" in w for w in warnings)


def test_corrupt_graph_json_is_survived(tmp_path):
    root = tmp_path / "m_BAD.DTEMPLATE"
    (root / "bin").mkdir(parents=True)
    (root / "bin" / "@3.bin").write_text("{not valid json")
    asset = Asset(name="m_BAD", obj_type="DTEMPLATE", path=root)
    warnings = []
    mapping = parse_mapping(asset, connections={}, warnings=warnings)
    assert mapping.transformations == []
    assert warnings


def test_graph_attachment_found_at_any_id(tmp_path):
    """Attachment ids are assigned by Informatica - never assume @3."""
    root = tmp_path / "m_ODD.DTEMPLATE"
    (root / "bin").mkdir(parents=True)
    (root / "bin" / "@17.bin").write_text(json.dumps({
        "content": {"transformations": [{"$$ID": 1, "$$class": 7, "name": "src"}],
                    "links": []},
        "metadata": {"$$classInfo": {
            "7": "com.informatica.metadata.template.tx.tmplsource.TmplSource"}},
    }))
    (root / "fileRecord.json").write_text(
        json.dumps([{"id": "@17", "type": "IMFOBJECT", "name": "m_ODD"}]))
    asset = Asset(name="m_ODD", obj_type="DTEMPLATE", path=root)
    mapping = parse_mapping(asset, connections={}, warnings=[])
    assert [t.name for t in mapping.transformations] == ["src"]


def test_unresolvable_connection_leaves_name_blank(tmp_path):
    """A connection outside the package must not break the mapping."""
    asset = _write_mapping(
        tmp_path, "m_X",
        content={"transformations": [{
            "$$ID": 1, "$$class": 7, "name": "src",
            "dataAdapter": {"connectionId": "saas:@NOT_IN_PACKAGE",
                            "typeSystem": "Oracle",
                            "object": {"objectName": "SOME_TABLE"}},
        }], "links": []},
        class_info={"7": "com.informatica.metadata.template.tx.tmplsource.TmplSource"},
    )
    mapping = parse_mapping(asset, connections={}, warnings=[])
    tx = mapping.transformations[0]
    assert tx.connection_name == ""
    assert tx.connection_type == "Oracle"      # falls back to the adapter's own type
    assert tx.object_name == "SOME_TABLE"


# --------------------------------------------------------------- rendering

def test_renderers_handle_an_empty_integration():
    """A package with nothing recognisable still produces valid output."""
    empty = Integration(meta=IntegrationMeta(taskflow_name="tf_EMPTY"))
    assert mapping_detail_rows(empty) == []
    assert field_level_rows(empty) == []


def test_rows_render_for_a_step_with_no_task():
    """Command and notification steps have no backing task asset."""
    integration = Integration(meta=IntegrationMeta(taskflow_name="tf_X"))
    integration.steps.append(
        Step(seq="1", title="Merge Files", step_type="Command Task")
    )
    rows = mapping_detail_rows(integration)
    assert len(rows) == 1
    assert rows[0][1] == "Merge Files"


def test_cyclic_mapping_links_do_not_hang():
    """Guard the flow walk against a cycle in the mapping graph."""
    mapping = Mapping(name="m_CYCLE")
    from iics_parser.model.ir import Transformation
    mapping.transformations = [
        Transformation(name="a", kind="Source"),
        Transformation(name="b", kind="Expression"),
    ]
    mapping.links = [{"from": "a", "to": "b"}, {"from": "b", "to": "a"}]
    assert flow_string(mapping)      # returns rather than recursing forever


def test_excel_and_word_render_a_minimal_integration(tmp_path):
    from iics_parser.render.excel import write_workbook
    from iics_parser.render.word import write_document

    integration = Integration(meta=IntegrationMeta(taskflow_name="tf_MIN"))
    integration.steps.append(Step(seq="1", title="Only Step", step_type="Command Task"))
    integration.connections.append(Connection(name="cn_x", conn_type="Oracle"))

    assert write_workbook(integration, tmp_path / "a.xlsx").exists()
    assert write_document(integration, tmp_path / "a.docx").exists()


# ------------------------------------------------- worksheet cell constraints

def test_control_characters_do_not_break_the_workbook(tmp_path):
    """SQL and email bodies carry control characters Excel's XML cannot hold."""
    from iics_parser.render.excel import write_workbook
    integration = Integration(meta=IntegrationMeta(taskflow_name="tf_CTRL"))
    step = Step(seq="1", title="Notify", step_type="Notification Task")
    step.parameters = [TaskParameter(name="Email_Body", value="line\x0bbreak\x00null")]
    integration.steps.append(step)

    path = write_workbook(integration, tmp_path / "ctrl.xlsx")
    import openpyxl
    values = [c.value for r in openpyxl.load_workbook(path)["Mapping Details"].iter_rows()
              for c in r if c.value]
    assert any("linebreaknull" in str(v) for v in values)


def test_oversized_cell_is_truncated_visibly(tmp_path):
    """Excel silently clips past 32,767 chars - say so rather than mislead."""
    from iics_parser.render.excel import MAX_CELL_CHARS, write_workbook
    integration = Integration(meta=IntegrationMeta(taskflow_name="tf_BIG"))
    step = Step(seq="1", title="Load", step_type="Mapping Task")
    step.parameters = [TaskParameter(name="Pre SQL", value="A" * (MAX_CELL_CHARS + 5000))]
    integration.steps.append(step)

    path = write_workbook(integration, tmp_path / "big.xlsx")
    import openpyxl
    wb = openpyxl.load_workbook(path)
    longest = max(len(str(c.value)) for r in wb["Mapping Details"].iter_rows()
                  for c in r if c.value)
    assert longest <= MAX_CELL_CHARS
    report = "\n".join(str(c.value) for r in wb["Parse Report"].iter_rows()
                       for c in r if c.value)
    assert "truncated" in report, "a shortened cell must be reported, not hidden"


def test_deep_mapping_chain_does_not_overflow_the_stack():
    """A long chain must not raise RecursionError."""
    from iics_parser.model.ir import Transformation
    mapping = Mapping(name="m_DEEP")
    n = 1500
    mapping.transformations = [Transformation(name=f"t{i}", kind="Expression")
                               for i in range(n)]
    mapping.links = [{"from": f"t{i}", "to": f"t{i+1}"} for i in range(n - 1)]
    assert flow_string(mapping)


def test_wide_mapping_does_not_produce_unbounded_paths():
    """A wide graph has combinatorially many paths; the output stays bounded."""
    from iics_parser.model.ir import Transformation
    from iics_parser.parse.mapping import _MAX_PATHS
    mapping = Mapping(name="m_WIDE")
    tx = [Transformation(name=f"src{i}", kind="Source") for i in range(12)]
    tx += [Transformation(name=f"mid{i}", kind="Expression") for i in range(12)]
    tx += [Transformation(name=f"tgt{i}", kind="Target") for i in range(12)]
    mapping.transformations = tx
    mapping.links = [{"from": f"src{i}", "to": f"mid{j}"}
                     for i in range(12) for j in range(12)]
    mapping.links += [{"from": f"mid{i}", "to": f"tgt{j}"}
                      for i in range(12) for j in range(12)]
    assert len(flow_string(mapping).splitlines()) <= _MAX_PATHS + 1


def test_extra_taskflows_are_reported(tmp_path):
    """Only one taskflow is analysed - a second must not vanish silently."""
    from iics_parser.extract.package import Asset, ExportPackage

    class TwoFlows(ExportPackage):
        def __init__(self):
            self.root = tmp_path
            self.metadata = {}
            self.warnings = []
            self.assets = [
                Asset(name="tf_A", obj_type="TASKFLOW", path=tmp_path / "a.xml"),
                Asset(name="tf_B", obj_type="TASKFLOW", path=tmp_path / "b.xml"),
            ]

    from iics_parser.parse.build import build_integration
    (tmp_path / "a.xml").write_text("<root/>")
    (tmp_path / "b.xml").write_text("<root/>")
    integration = build_integration(TwoFlows())
    assert any("2 taskflows" in w for w in integration.warnings)
