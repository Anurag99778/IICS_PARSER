"""A package with no taskflow must still be documented.

IICS exports a single asset as readily as a whole orchestration, so plenty of
the integrations being migrated are one mapping, or a mapping task and the
mapping it runs. Every renderer works from ``integration.steps``, and steps
used to come only from a taskflow - so these packages produced a workbook with
three empty sheets and a Word document with nothing in it.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from iics_parser.extract.package import ExportPackage
from iics_parser.parse.build import build_integration
from iics_parser.pipeline import process
from iics_parser.render.rows import (field_level_rows, mapping_detail_rows,
                                     object_field_rows)

from . import synthetic


@pytest.fixture()
def with_task(tmp_path):
    zip_path = synthetic.write_package(
        tmp_path / "one_mapping.zip", task_name="mct_LOAD_EMPLOYEES")
    return build_integration(ExportPackage.open(zip_path, tmp_path / "x1"),
                             image_dir=tmp_path / "img1")


@pytest.fixture()
def bare(tmp_path):
    """A mapping exported entirely on its own - no task, no taskflow."""
    zip_path = synthetic.write_package(tmp_path / "bare.zip")
    return build_integration(ExportPackage.open(zip_path, tmp_path / "x2"),
                             image_dir=tmp_path / "img2")


# ------------------------------------------------------- a mapping and a task

def test_a_mapping_task_without_a_taskflow_becomes_a_step(with_task):
    assert len(with_task.steps) == 1
    step = with_task.steps[0]
    assert step.seq == "1"
    assert step.step_type == "Mapping Task"
    assert step.task_name == "mct_LOAD_EMPLOYEES"
    assert step.mapping_name == "m_LOAD_EMPLOYEES"


def test_its_rows_reach_every_sheet(with_task):
    detail = mapping_detail_rows(with_task)
    fields = field_level_rows(with_task)
    objects = object_field_rows(with_task)
    assert detail and fields and objects

    blob = "\n".join(str(c) for rows in (detail, fields, objects)
                     for r in rows for c in r)
    assert "STG_EMPLOYEES" in blob                    # source object
    assert "DW_EMPLOYEES" in blob                     # target object
    assert "cn_WAREHOUSE (Oracle)" in blob            # connection
    assert "UPPER(FULL_NAME)" in blob                 # expression
    assert "→ FULL_NAME" in blob                      # column lineage
    assert "src_employees->exp_clean->tgt_employees" in blob


def test_task_parameters_are_documented(with_task):
    blob = "\n".join(str(c) for r in mapping_detail_rows(with_task) for c in r)
    assert "employees.param" in blob
    assert "p_run_date" in blob


# --------------------------------------------------------- a bare mapping

def test_a_bare_mapping_is_documented_too(bare):
    assert len(bare.steps) == 1
    step = bare.steps[0]
    assert step.step_type == "Mapping"
    assert step.task_name == ""            # there genuinely is no task
    assert step.mapping_name == "m_LOAD_EMPLOYEES"
    assert bare.tasks == []                # and none is invented

    blob = "\n".join(str(c) for r in mapping_detail_rows(bare) for c in r)
    assert "STG_EMPLOYEES" in blob and "DW_EMPLOYEES" in blob


def test_the_package_is_named_after_what_was_exported(bare, with_task):
    assert bare.meta.taskflow_name == "m_LOAD_EMPLOYEES"
    assert with_task.meta.taskflow_name == "mct_LOAD_EMPLOYEES"
    # The timestamp IICS appends to the export name is not part of the name.
    assert "1787142080634" not in bare.meta.taskflow_name


def test_a_single_asset_lends_the_package_its_description(bare):
    """One asset's description is the whole of what this package does."""
    assert bare.meta.description == "Load employees into the warehouse."


def test_the_absence_of_execution_order_is_stated(bare):
    """Name order is not execution order, and a reader must not assume it is."""
    assert any("no taskflow" in w and "name order" in w for w in bare.warnings), \
        bare.warnings


# ------------------------------------------------------------- end to end

def test_both_documents_are_written_and_are_not_empty(tmp_path):
    zip_path = synthetic.write_package(tmp_path / "one.zip",
                                       task_name="mct_LOAD_EMPLOYEES")
    result = process(zip_path, tmp_path / "out")

    assert result.ok, result.error
    assert result.excel.exists() and result.word.exists()
    assert result.excel.stem.startswith("mct_LOAD_EMPLOYEES")

    from openpyxl import load_workbook
    wb = load_workbook(result.excel)
    for sheet in ("Mapping Details", "Field Values", "Source & Target Fields"):
        rows = [r for r in wb[sheet].iter_rows(min_row=5, values_only=True)
                if any(v not in (None, "") for v in r)]
        assert rows, f"'{sheet}' came out empty"


def test_coverage_is_complete_for_a_single_mapping_package(with_task):
    from iics_parser import coverage
    lines = coverage.audit(with_task, mapping_detail_rows(with_task),
                           field_level_rows(with_task),
                           object_field_rows(with_task))
    incomplete = [(l.category, l.summary, l.missing) for l in lines if not l.complete]
    assert not incomplete, f"objects missing from the output: {incomplete}"
