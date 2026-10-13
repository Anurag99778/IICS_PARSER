"""The branded front page, the AI purpose section, and Connection Details.

The template carries a cover page in a section of its own. A section is defined
by a break on the *last paragraph of that section*, so clearing every block -
which the renderer used to do - deleted the section, and with it the cover
artwork, which lives in that section's own header.
"""

from __future__ import annotations

import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook

from iics_parser.extract.package import ExportPackage
from iics_parser.parse.build import build_integration
from iics_parser.pipeline import process
from iics_parser.render.rows import (CONNECTION_DETAIL_COLUMNS,
                                     connection_detail_rows)
from iics_parser.render.word import TEMPLATE, write_document

from . import synthetic

PURPOSE = ("Loads team member demographic data from the Cloud Integration Hub "
           "into the WFM staging schema each night.")


@pytest.fixture()
def integration(tmp_path):
    zip_path = synthetic.write_package(
        tmp_path / "pkg.zip", project="WFM",
        mapping_name="m_TEAM_MEMBER_LOAD", task_name="mt_TEAM_MEMBER_LOAD",
        description=PURPOSE)
    return build_integration(ExportPackage.open(zip_path, tmp_path / "x"),
                             image_dir=tmp_path / "img")


@pytest.fixture()
def document(integration, tmp_path):
    return write_document(integration, tmp_path / "out.docx")


# --------------------------------------------------------------- front page

@pytest.mark.skipif(not TEMPLATE.exists(), reason="template not installed")
def test_the_cover_page_survives(document):
    """Two sections means the cover section was not collapsed into the body."""
    assert len(Document(document).sections) == 2

    with zipfile.ZipFile(document) as docx:
        names = docx.namelist()
        assert sum(1 for n in names if n.startswith("word/header")) >= 2
        assert sum(1 for n in names if "media/" in n) > 5, "cover artwork lost"


@pytest.mark.skipif(not TEMPLATE.exists(), reason="template not installed")
def test_the_front_tables_carry_this_integration_not_the_templates(document):
    tables = Document(document).tables
    details = {r.cells[0].text.strip(): r.cells[1].text.strip()
               for r in tables[0].rows}

    assert details["Integration Name"] == "mt_TEAM_MEMBER_LOAD"
    assert details["Document Type"] == "Process Document"
    assert details["OIC Architecture"] == "Oracle Integration Cloud"
    # "Business Process Name" also starts with "Business Process"; the longer
    # label has to win or it inherits the wrong value.
    assert details["Business Process"] == "IICS to OIC Migration"
    assert details["Business Process Name"] == "[to be confirmed]"


@pytest.mark.skipif(not TEMPLATE.exists(), reason="template not installed")
def test_no_content_from_the_template_document_is_left_behind(document):
    """The template ships with another integration's body - it must not leak."""
    text = "\n".join(p.text for p in Document(document).paragraphs)
    assert "CONCUR" not in text.upper()


@pytest.mark.skipif(not TEMPLATE.exists(), reason="template not installed")
def test_word_is_asked_to_rebuild_the_contents_page(document):
    """The contents block is kept from the template, so it must be refreshed."""
    with zipfile.ZipFile(document) as docx:
        assert b"updateFields" in docx.read("word/settings.xml")


def test_a_plain_template_still_gets_a_front_page(integration, tmp_path):
    """With no cover section to keep, the renderer builds the tables itself."""
    plain = tmp_path / "plain.docx"
    Document().save(plain)

    path = write_document(integration, tmp_path / "plain_out.docx", template=plain)
    tables = Document(path).tables
    labels = [r.cells[0].text.strip() for r in tables[0].rows]
    assert "Integration Name" in labels


# ------------------------------------------------- integration purpose

def test_the_purpose_section_states_this_integrations_purpose(document):
    """The AI/override draft, not only the standing migration objective."""
    text = [p.text.strip() for p in Document(document).paragraphs]
    start = text.index("Integration Purpose/Objective:")
    section = [line for line in text[start + 1:start + 5] if line]

    assert section[0] == PURPOSE
    assert any("migration of all IICS" in line for line in section)


def test_an_undrafted_purpose_is_visibly_missing_not_papered_over(tmp_path):
    """Boilerplate must never stand in for the purpose nobody supplied."""
    zip_path = synthetic.write_package(tmp_path / "nodesc.zip",
                                       task_name="mct_X", description="")
    integration = build_integration(ExportPackage.open(zip_path, tmp_path / "x"))
    path = write_document(integration, tmp_path / "nodesc.docx")

    text = [p.text.strip() for p in Document(path).paragraphs]
    start = text.index("Integration Purpose/Objective:")
    assert "[to be confirmed]" in text[start + 1:start + 4]


# ------------------------------------------------------ connection details

def test_connection_details_rows_stand_alone(integration):
    rows = connection_detail_rows(integration)
    assert len(rows) == 1
    row = dict(zip(CONNECTION_DETAIL_COLUMNS, rows[0]))

    assert row["Schema"] == "APP_OWNER"
    assert row["Source Connection"] == "cn_WAREHOUSE (Oracle)"
    assert row["Source Object"] == "STG_EMPLOYEES"
    assert row["Target Connection"] == "cn_WAREHOUSE (Oracle)"
    assert row["Target Object"] == "DW_EMPLOYEES"
    assert row["Mapping Name"] == "m_TEAM_MEMBER_LOAD"


def test_the_sheet_is_in_the_workbook(integration, tmp_path):
    from iics_parser.render.excel import CONNECTION_SHEET, write_workbook
    workbook = load_workbook(write_workbook(integration, tmp_path / "a.xlsx"))

    assert CONNECTION_SHEET in workbook.sheetnames
    sheet = workbook[CONNECTION_SHEET]
    assert [c.value for c in sheet[4]] == CONNECTION_DETAIL_COLUMNS
    assert sheet["A5"].value == "1"


def test_a_schema_is_taken_from_the_object_or_the_connection(integration):
    """dbSchema when the export names one; otherwise the connection's own."""
    from iics_parser.render.rows import _schema_label

    source, target = integration.mappings[0].sources[0], integration.mappings[0].targets[0]
    schemas = {c.name: c.schema_label for c in integration.connections}

    assert _schema_label(schemas, source, target) == "APP_OWNER"

    source.db_schema = "RAW"          # both ends named, and they differ
    assert _schema_label(schemas, source, target) == "RAW → APP_OWNER"

    assert _schema_label(schemas, None, None) == ""


def test_documents_are_written_end_to_end(tmp_path):
    zip_path = synthetic.write_package(tmp_path / "e2e.zip", task_name="mct_X")
    result = process(zip_path, tmp_path / "out")

    assert result.ok, result.error
    workbook = load_workbook(result.excel)
    rows = [r for r in workbook["Connection Details"].iter_rows(
        min_row=5, values_only=True) if any(v not in (None, "") for v in r)]
    assert rows


# ------------------------------------------------------------- cover page

def _cover_paragraphs(path):
    from docx.oxml.ns import qn
    document = Document(path)
    return [("".join(t.text or "" for t in p.findall(qn("w:t")) or []) or
             "".join(t.text or "" for t in p.iter(qn("w:t"))))
            for p in document.element.body.iter(qn("w:p"))]


def _title_lines(path):
    """The runs of each cover-title paragraph, as (size, text)."""
    from docx.oxml.ns import qn
    out = []
    for paragraph in Document(path).element.body.iter(qn("w:p")):
        runs = paragraph.findall(qn("w:r"))
        texts = ["".join(t.text or "" for t in r.findall(qn("w:t"))) for r in runs]
        if "IICS to OIC Modernization" in "".join(texts):
            out.append(texts)
    return out


@pytest.mark.skipif(not TEMPLATE.exists(), reason="template not installed")
def test_the_cover_names_the_integration(document):
    """Otherwise every integration arrives with an identical front page."""
    titles = _title_lines(document)
    assert titles, "cover title paragraph not found"

    # Word stores the cover twice - DrawingML and a VML fallback - and both
    # have to say the same thing or the two renderings disagree.
    assert len(titles) == 2
    for runs in titles:
        assert runs[0].strip() == "IICS to OIC Modernization"
        assert "mt_TEAM_MEMBER_LOAD" in runs[-1]


@pytest.mark.skipif(not TEMPLATE.exists(), reason="template not installed")
def test_the_name_is_on_its_own_line_and_not_title_sized(document):
    from docx.oxml.ns import qn

    for paragraph in Document(document).element.body.iter(qn("w:p")):
        runs = paragraph.findall(qn("w:r"))
        texts = ["".join(t.text or "" for t in r.findall(qn("w:t"))) for r in runs]
        if "mt_TEAM_MEMBER_LOAD" not in "".join(texts):
            continue
        title_run, name_run = runs[0], runs[-1]
        assert name_run.find(qn("w:br")) is not None, "name runs on after the title"

        def size(run):
            node = run.find(qn("w:rPr") + "/" + qn("w:sz"))
            return int(node.get(qn("w:val"))) if node is not None else None

        assert size(name_run) < size(title_run), "a long name at title size overflows"
        return
    pytest.fail("no cover paragraph carried the integration name")


@pytest.mark.skipif(not TEMPLATE.exists(), reason="template not installed")
def test_the_cover_is_dated_today_not_when_the_template_was_written(document):
    from datetime import date
    today = date.today()

    dates = [p for p in _cover_paragraphs(document)
             if p.strip().endswith(f"{today:%B %Y}")]
    assert dates, "the template's own submission date was left in place"
    assert all(str(today.day) in p for p in dates)
    assert not any("19th August 2026" in p for p in _cover_paragraphs(document))


@pytest.mark.skipif(not TEMPLATE.exists(), reason="template not installed")
def test_generating_twice_does_not_stack_the_name_up(integration, tmp_path):
    """Each run starts from the template, so nothing may accumulate."""
    first = write_document(integration, tmp_path / "one.docx")
    second = write_document(integration, tmp_path / "two.docx")

    for path in (first, second):
        for runs in _title_lines(path):
            assert "".join(runs).count("mt_TEAM_MEMBER_LOAD") == 1


def test_a_template_with_no_cover_title_is_left_alone(integration, tmp_path):
    """An unfamiliar template keeps its own wording rather than being mangled."""
    plain = tmp_path / "plain.docx"
    blank = Document()
    blank.add_paragraph("Some Other Company Template")
    blank.save(plain)

    path = write_document(integration, tmp_path / "plain_out.docx", template=plain)
    assert path.exists()


def test_a_nameless_integration_does_not_write_an_empty_line(tmp_path):
    """Guard against decorating the cover with nothing."""
    zip_path = synthetic.write_package(tmp_path / "x.zip", task_name="mct_X")
    built = build_integration(ExportPackage.open(zip_path, tmp_path / "x"))
    built.meta.taskflow_name = built.meta.display_name = ""

    path = write_document(built, tmp_path / "nameless.docx")
    for runs in _title_lines(path):
        assert len(runs) == 1, "an empty second line was added to the title"
